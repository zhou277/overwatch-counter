from pathlib import Path
import json
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data.xlsx"
TEMPLATE = ROOT / "template.html"
OUTPUT = ROOT / "index.html"
ROLE_CN = {"T":"坦克","C":"输出","S":"支援"}

def clean(v):
    return "" if v is None else str(v).strip()

def rows_as_dicts(ws):
    rows=list(ws.iter_rows(values_only=True))
    if not rows: return []
    headers=[clean(x) for x in rows[0]]
    out=[]
    for row in rows[1:]:
        if all(clean(v)=="" for v in row): continue
        out.append({headers[i]: row[i] for i in range(len(headers))})
    return out

def tip_card(rank,text,top10):
    css="top10" if top10 else "secondary"
    return f'<div class="tip-card {css}"><div class="tip-rank">Top {rank}</div><div class="tip-text">{text}</div></div>'

def main():
    wb=load_workbook(DATA,data_only=True)
    settings={clean(r["字段"]):clean(r["值"]) for r in rows_as_dicts(wb["settings"])}

    hero_rows=sorted(rows_as_dicts(wb["heroes"]),key=lambda r:int(r["order"] or 9999))
    heroes=[]
    for r in hero_rows:
        heroes.append({
            "name":clean(r["name"]),
            "aliases":[x.strip() for x in clean(r["aliases"]).split(",") if x.strip()],
            "role":ROLE_CN.get(clean(r["role"]).upper(),clean(r["role"])),
            "strong":[],
            "weak":[],
            "note":clean(r["note"])
        })

    hero_names=[h["name"] for h in heroes]
    matchup={n:{"strong":{"坦克":[],"输出":[],"支援":[]},"weak":{"坦克":[],"输出":[],"支援":[]}} for n in hero_names}
    reason={n:{"strong":{"坦克":[],"输出":[],"支援":[]},"weak":{"坦克":[],"输出":[],"支援":[]}} for n in hero_names}

    counters=rows_as_dicts(wb["counters"])
    counters.sort(key=lambda r:(
        hero_names.index(clean(r["hero"])) if clean(r["hero"]) in hero_names else 9999,
        clean(r["relation"]),
        {"T":0,"C":1,"S":2}.get(clean(r["target_role"]).upper(),9),
        int(r["rank"] or 999)
    ))
    for r in counters:
        hero=clean(r["hero"]); relation=clean(r["relation"]).lower()
        role=ROLE_CN.get(clean(r["target_role"]).upper(),clean(r["target_role"]))
        target=clean(r["target"]); why=clean(r["reason"])
        if hero not in matchup or relation not in ("strong","weak") or role not in ("坦克","输出","支援") or not target:
            continue
        matchup[hero][relation][role].append(target)
        reason[hero][relation][role].append({"name":target,"reason":why})

    by_name={h["name"]:h for h in heroes}
    for n,data in matchup.items():
        by_name[n]["strong"]=sum(data["strong"].values(),[])
        by_name[n]["weak"]=sum(data["weak"].values(),[])

    map_rows=sorted(rows_as_dicts(wb["maps"]),key=lambda r:int(r["order"] or 9999))
    maps=[]
    for r in map_rows:
        maps.append({
            "name":clean(r["name"]),"mode":clean(r["mode"]),"style":clean(r["style"]),
            "tank":[clean(r[f"tank{i}"]) for i in range(1,4) if clean(r[f"tank{i}"])],
            "dps":[clean(r[f"dps{i}"]) for i in range(1,4) if clean(r[f"dps{i}"])],
            "support":[clean(r[f"support{i}"]) for i in range(1,4) if clean(r[f"support{i}"])],
            "comp":clean(r["composition"]),"reason":clean(r["reason"])
        })

    tips=sorted(rows_as_dicts(wb["tips"]),key=lambda r:int(r["rank"] or 9999))
    template=TEMPLATE.read_text(encoding="utf-8")
    output=template
    output=output.replace("__HEROES_JSON__",json.dumps(heroes,ensure_ascii=False,separators=(",",":")))
    output=output.replace("__MATCHUP_JSON__",json.dumps(matchup,ensure_ascii=False,separators=(",",":")))
    output=output.replace("__REASON_JSON__",json.dumps(reason,ensure_ascii=False,separators=(",",":")))
    output=output.replace("__MAPS_JSON__",json.dumps(maps,ensure_ascii=False,separators=(",",":")))
    output=output.replace("__TOP10_HTML__","\n".join(tip_card(int(t["rank"]),clean(t["text"]),True) for t in tips if int(t["rank"])<=10))
    output=output.replace("__TOPREST_HTML__","\n".join(tip_card(int(t["rank"]),clean(t["text"]),False) for t in tips if int(t["rank"])>10))
    output=output.replace("__VERSION__",settings.get("version","v3.1.4"))
    OUTPUT.write_text(output,encoding="utf-8")
    print(f"Generated {OUTPUT.name}: {len(heroes)} heroes, {len(counters)} counter rows, {len(maps)} maps, {len(tips)} tips")

if __name__=="__main__":
    main()
